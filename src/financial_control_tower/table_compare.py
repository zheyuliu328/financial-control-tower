"""Offline, value-only comparison of two explicitly mapped tables."""

import csv
import ctypes
import hashlib
import html
import io
import json
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Optional
from zipfile import ZipFile

from . import __version__

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_FIELD_CHECKS = 20_000
NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?\Z")
BLOCKED = {
    "missing_key",
    "invalid_key",
    "duplicate_key",
    "invalid_numeric",
    "formula_cell",
    "invalid_currency",
    "currency_mismatch",
}
STATUS_LABELS = {
    "matched": "Within tolerance",
    "amount_mismatch": "Amount difference",
    "left_only": "Only in left",
    "right_only": "Only in right",
    "missing_key": "Missing key",
    "invalid_key": "Invalid key",
    "duplicate_key": "Duplicate key",
    "invalid_numeric": "Invalid numeric value",
    "formula_cell": "Excel formula",
    "invalid_currency": "Invalid currency",
    "currency_mismatch": "Currency mismatch",
}


@dataclass(frozen=True)
class TableSpec:
    """Column names are exact and mappings are positional across the two sides."""

    path: Path
    keys: tuple
    values: tuple
    sheet: Optional[str] = None
    header_row: int = 1
    currency: Optional[str] = None


def _number(value):
    text = str(value).strip()
    if len(text) > 128 or not NUMBER.fullmatch(text):
        raise ValueError("expected a finite decimal number without grouping separators")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("invalid decimal exponent") from exc
    if len(value.as_tuple().digits) > 50 or not -100 <= value.adjusted() <= 100:
        raise ValueError("number exceeds the supported 50-digit / exponent -100..100 range")
    return value


def _snapshot(path):
    path = Path(path).absolute()
    before = path.stat()
    if not path.is_file() or before.st_size > MAX_BYTES:
        raise ValueError(f"input must be a regular file of at most 25 MiB: {path}")
    with path.open("rb") as stream:
        content = stream.read(MAX_BYTES + 1)
    after = path.stat()
    fields = ("st_dev", "st_ino", "st_size", "st_mtime_ns", "st_ctime_ns")
    identity = tuple(getattr(after, field) for field in fields)
    if len(content) > MAX_BYTES or identity != tuple(getattr(before, field) for field in fields):
        raise ValueError(f"input changed while being read; use a stable copy: {path}")
    return content, identity, hashlib.sha256(content).hexdigest()


def _display(value):
    return "" if value is None else str(value)


def _raw_rows(spec, content, metadata):
    """Return row values and cell kinds without evaluating Excel formulas."""
    if spec.path.suffix.lower() == ".csv":
        if spec.sheet is not None:
            raise ValueError("a sheet name applies only to .xlsx inputs")
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        for row in reader:
            yield [(value, "text") for value in row]
        return
    if spec.path.suffix.lower() != ".xlsx":
        raise ValueError("supported input formats are UTF-8 comma-separated .csv and value-only .xlsx")
    with ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if (
            len(members) > 2048
            or sum(member.file_size for member in members) > 100 * 1024 * 1024
            or any(member.file_size > MAX_BYTES for member in members)
        ):
            raise ValueError("Excel archive exceeds the 25 MiB member / 100 MiB total expanded / 2,048 member limit")
    try:
        from defusedxml.common import DefusedXmlException
        from defusedxml.ElementTree import ParseError
        from openpyxl import LXML, load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise ValueError(
            "Excel input requires the optional extra: pip install 'financial-control-tower[excel]'"
        ) from exc
    parser_errors = (ParseError, DefusedXmlException, InvalidFileException, KeyError, IndexError, TypeError, ValueError)
    if LXML:
        from lxml.etree import XMLSyntaxError

        parser_errors += (XMLSyntaxError,)
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if spec.sheet is None and len(workbook.sheetnames) != 1:
            raise ValueError("select an explicit sheet when an Excel workbook has multiple sheets")
        sheet_name = spec.sheet if spec.sheet is not None else workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Excel sheet does not exist: {sheet_name}")
        metadata["sheet"] = sheet_name
        sheet = workbook[sheet_name]
        if not hasattr(sheet, "iter_rows") or not hasattr(sheet, "reset_dimensions"):
            raise ValueError("select a worksheet with tabular cells; chart sheets cannot be compared")
        # Do not trust an incorrect saved dimension to truncate non-empty cells.
        sheet.reset_dimensions()
        for row in sheet.iter_rows():
            yield [
                (
                    cell.value,
                    "formula"
                    if cell.data_type == "f"
                    else "text"
                    if isinstance(cell.value, str) and cell.data_type != "e"
                    else "other",
                )
                for cell in row
            ]
    except parser_errors as exc:
        # These exceptions are normalized only at the OOXML loader/row-reader boundary.
        # Bad ZIP member references, malformed XML, and invalid typed workbook fields
        # are input failures, not a completed comparison with differences (CLI exit 1).
        raise ValueError(f"invalid Excel workbook structure ({type(exc).__name__}): {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _read_table(spec, content):
    rows, blank_rows, headers = [], [], None
    metadata = {}
    needed = set(spec.keys + spec.values + ((spec.currency,) if spec.currency else ()))
    for row_number, cells in enumerate(_raw_rows(spec, content, metadata), 1):
        if row_number > MAX_ROWS + spec.header_row:
            raise ValueError("input exceeds the 100,000 data-row limit")
        if len(cells) > MAX_COLUMNS:
            raise ValueError("input exceeds the 200-column limit")
        if row_number < spec.header_row:
            continue
        if row_number == spec.header_row:
            # Excel may include trailing formatting-only cells, but internal blank headers are ambiguous.
            if spec.path.suffix.lower() == ".xlsx":
                while cells and cells[-1][0] is None:
                    cells.pop()
            if not cells or any(kind != "text" or not value.strip() for value, kind in cells):
                raise ValueError("header cells must be nonempty text; formulas and blank headers are rejected")
            headers = [value for value, _kind in cells]
            if len(set(headers)) != len(headers):
                raise ValueError("duplicate column headers are rejected before parsing")
            if missing := needed - set(headers):
                raise ValueError(f"mapped columns absent from the header: {sorted(missing)}")
            continue
        if not any(value is not None and value != "" for value, _kind in cells):
            blank_rows.append(row_number)
            continue
        if len(cells) != len(headers):
            if spec.path.suffix.lower() == ".csv" or any(value is not None for value, _kind in cells[len(headers) :]):
                raise ValueError(f"row {row_number} has a different width from its header")
            cells = cells[: len(headers)] + [(None, "other")] * max(0, len(headers) - len(cells))
        selected = {name: cells[index] for index, name in enumerate(headers) if name in needed}
        issues = []
        key = []
        for name in spec.keys:
            value, kind = selected[name]
            if kind == "formula":
                issues.append(("formula_cell", f"{name}: Excel formula; provide independently checked, frozen values"))
            elif value is None or not _display(value).strip():
                issues.append(("missing_key", f"{name}: empty key"))
            elif kind != "text":
                issues.append(("invalid_key", f"{name}: Excel keys must be stored as text to preserve leading zeros"))
            key.append(_display(value))
        numbers = {}
        for name in spec.values:
            value, kind = selected[name]
            if kind == "formula":
                issues.append(("formula_cell", f"{name}: Excel formula; cached results do not prove freshness"))
                continue
            try:
                numbers[name] = _number(value)
            except ValueError as exc:
                issues.append(("invalid_numeric", f"{name}: {exc}"))
        currency = ""
        if spec.currency:
            value, kind = selected[spec.currency]
            currency = _display(value)
            if kind == "formula":
                issues.append(("formula_cell", f"{spec.currency}: Excel formula in currency"))
            elif kind != "text" or not currency.strip():
                issues.append(("invalid_currency", f"{spec.currency}: currency must be nonempty text"))
        invalid_key = any(kind in {"missing_key", "invalid_key"} for kind, _detail in issues) or any(
            selected[name][1] == "formula" for name in spec.keys
        )
        rows.append(
            {
                "row": row_number,
                "key": tuple(key),
                "invalid_key": invalid_key,
                "raw": {name: _display(value) for name, (value, _kind) in selected.items()},
                "numbers": numbers,
                "currency": currency,
                "issues": issues,
            }
        )
    if headers is None:
        raise ValueError("selected header row does not exist")
    if not rows:
        raise ValueError("table contains no nonempty data rows; an empty comparison cannot pass")
    return rows, {**metadata, "headers": headers, "data_rows": len(rows), "blank_rows_ignored": blank_rows}


def _validate_specs(left, right, absolute_tolerance, relative_tolerance, common_unit):
    for spec in (left, right):
        if not spec.keys or not spec.values:
            raise ValueError("at least one key column and one value column are required on both sides")
        if any(not isinstance(name, str) or not name.strip() for name in spec.keys + spec.values):
            raise ValueError("mapped column names must be nonempty strings")
        if len(set(spec.keys)) != len(spec.keys) or len(set(spec.values)) != len(spec.values):
            raise ValueError("a key/value mapping cannot repeat a column on the same side")
        if type(spec.header_row) is not int or spec.header_row < 1:
            raise ValueError("header row must be a positive integer")
    if len(left.keys) != len(right.keys) or len(left.values) != len(right.values):
        raise ValueError("left/right key mappings and value mappings must have equal lengths")
    if bool(left.currency) != bool(right.currency):
        raise ValueError("currency columns must be specified on both sides")
    if left.currency:
        if common_unit is not None:
            raise ValueError("choose currency columns or a common-unit declaration, not both")
    elif not isinstance(common_unit, str) or not common_unit.strip():
        raise ValueError("without currency columns, explicitly declare --common-unit (for example USD)")
    absolute, relative = _number(absolute_tolerance), _number(relative_tolerance)
    if absolute < 0 or relative < 0:
        raise ValueError("tolerances must be nonnegative; relative tolerance is a fraction, not percent")
    return absolute, relative


def _compare_rows(left_rows, right_rows, left, right, absolute, relative):
    records, groups = [], defaultdict(lambda: {"left": [], "right": []})

    def emit(a, b, forced=None):
        issues = [
            f"{side} row {row['row']}: {detail}"
            for side, row in (("left", a), ("right", b))
            if row
            for _kind, detail in row["issues"]
        ]
        status = forced
        if status is None:
            errors = [kind for row in (a, b) if row for kind, _detail in row["issues"]]
            if errors:
                status = errors[0]
            elif a is None or b is None:
                status = "right_only" if a is None else "left_only"
                absent_side = "left" if a is None else "right"
                issues.append(
                    f"No corresponding key in the {absent_side} file. Check extract coverage and key-column mapping."
                )
            elif a["currency"] != b["currency"]:
                status = "currency_mismatch"
                issues.append("currency text differs; no subtraction or FX conversion performed")
        if forced == "duplicate_key":
            issues.append("key has multiple rows on at least one side; every member is retained without pairing")
        for left_name, right_name in zip(left.values, right.values):
            difference = allowed = None
            field_status = status
            field_issues = list(issues)
            if status is None:
                with localcontext() as context:
                    context.prec = 500
                    av, bv = a["numbers"][left_name], b["numbers"][right_name]
                    difference = bv - av
                    allowed = max(absolute, relative * max(abs(av), abs(bv)))
                    field_status = "matched" if abs(difference) <= allowed else "amount_mismatch"
                    if field_status == "amount_mismatch":
                        field_issues.append(
                            f"Absolute difference {difference.copy_abs()} exceeds the selected allowance {allowed}. Check source amounts, units and field mapping; the comparison does not determine the cause."
                        )
            records.append(
                {
                    "status": field_status,
                    "key": list((a or b)["key"]),
                    "left_row": a["row"] if a else None,
                    "right_row": b["row"] if b else None,
                    "left_column": left_name,
                    "right_column": right_name,
                    "left_value": a["raw"][left_name] if a else None,
                    "right_value": b["raw"][right_name] if b else None,
                    "left_currency": a["currency"] if a else None,
                    "right_currency": b["currency"] if b else None,
                    "difference_right_minus_left": str(difference) if difference is not None else None,
                    "absolute_difference": str(difference.copy_abs()) if difference is not None else None,
                    "allowed_difference": str(allowed) if allowed is not None else None,
                    "issues": field_issues,
                }
            )

    for side, rows in (("left", left_rows), ("right", right_rows)):
        for row in rows:
            if row["invalid_key"]:
                emit(row if side == "left" else None, row if side == "right" else None)
            else:
                groups[row["key"]][side].append(row)
    for group in groups.values():
        a, b = group["left"], group["right"]
        if len(a) > 1 or len(b) > 1:
            for row in a:
                emit(row, None, "duplicate_key")
            for row in b:
                emit(None, row, "duplicate_key")
        else:
            emit(a[0] if a else None, b[0] if b else None)
    return records


def _csv_text(value):
    """Quote dangerous spreadsheet text; comparison.json preserves the exact originals."""
    text = "" if value is None else str(value)
    probe = text.lstrip().lstrip("\ufeff").lstrip()
    if (probe and probe[0] in "=+-@") or any(ord(char) < 32 or ord(char) == 127 for char in text):
        return "'" + text
    return text


def _write_csv(path, records):
    columns = [
        "status",
        "key",
        "left_row",
        "right_row",
        "left_column",
        "right_column",
        "left_currency",
        "right_currency",
        "left_value",
        "right_value",
        "difference_right_minus_left",
        "absolute_difference",
        "allowed_difference",
        "issues",
    ]
    numeric = {"left_row", "right_row", "difference_right_minus_left", "absolute_difference", "allowed_difference"}
    with path.open("x", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(columns)
        for record in records:
            row = []
            for name in columns:
                value = record[name]
                if isinstance(value, list):
                    value = json.dumps(value, ensure_ascii=False)
                row.append(value if name in numeric else _csv_text(value))
            writer.writerow(row)


def _write_html(path, result):
    escape = lambda value: html.escape("" if value is None else str(value), quote=True)  # noqa: E731
    summary = result["summary"]
    counts = summary["field_checks"]
    cards = [
        ("Matched fields", counts.get("matched", 0)),
        ("Different fields", counts.get("amount_mismatch", 0)),
        ("Unpaired fields", sum(counts.get(name, 0) for name in ("left_only", "right_only"))),
        ("Blocked fields", sum(counts.get(name, 0) for name in BLOCKED)),
    ]
    cards_html = "".join(
        f"<div class='card'><strong>{count}</strong><span>{label}</span></div>" for label, count in cards
    )
    inputs_html = "".join(
        f"<li><strong>{side.title()}</strong>: {escape(item['path'])}<br>"
        f"Header row {item['header_row']}; sheet {escape(item['sheet'] or '(CSV / single worksheet)')}; "
        f"{item['data_rows']} data rows; {len(item['blank_rows_ignored'])} blank rows recorded as ignored.</li>"
        for side, item in result["inputs"].items()
    )
    row_html = []
    for record in result["records"]:
        state = "ok" if record["status"] == "matched" else "issue"
        values = [
            STATUS_LABELS[record["status"]],
            json.dumps(record["key"], ensure_ascii=False),
            record["left_row"],
            record["right_row"],
            f"{record['left_column']} / {record['right_column']}",
            record["left_currency"],
            record["right_currency"],
            record["left_value"],
            record["right_value"],
            record["difference_right_minus_left"],
            record["allowed_difference"],
            "; ".join(record["issues"]),
        ]
        row_html.append(f"<tr class='{state}'>" + "".join(f"<td>{escape(value)}</td>" for value in values) + "</tr>")
    headings = [
        "Status",
        "Exact composite key",
        "Left row",
        "Right row",
        "Left / right field",
        "Left currency",
        "Right currency",
        "Left value",
        "Right value",
        "Right − left",
        "Allowed difference",
        "Explanation",
    ]
    title = "All compared fields agree" if summary["all_matched"] else "Comparison completed with exceptions"
    policy = result["policy"]
    document = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'">
<title>Table comparison</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f4f6f9;color:#172637;font:15px/1.55 system-ui,sans-serif}}
main{{max-width:1600px;margin:auto;padding:36px 24px}}h1{{font-size:30px;margin:8px 0}}h2{{font-size:20px;margin-top:30px}}
.eyebrow{{font-size:12px;letter-spacing:.12em;text-transform:uppercase;color:#4b6074}}.intro{{max-width:880px}}
.cards{{display:grid;grid-template-columns:repeat(4,minmax(130px,1fr));gap:14px;margin:25px 0}}
.card{{background:white;border:1px solid #dce3eb;border-radius:9px;padding:18px}}.card strong{{display:block;font-size:30px}}
.card span{{color:#465d73}}a{{color:#164f95}}.links{{display:flex;gap:20px;flex-wrap:wrap}}li{{margin:12px 0;overflow-wrap:anywhere}}
.table-wrap{{overflow:auto;border:1px solid #dce3eb;border-radius:8px;background:white}}table{{border-collapse:collapse;width:100%;font-size:13px}}
th,td{{padding:11px 12px;border-bottom:1px solid #e3e8ef;text-align:left;vertical-align:top}}th{{background:#eaf0f6;white-space:nowrap}}
td{{white-space:pre-wrap;min-width:90px;max-width:380px;overflow-wrap:anywhere}}.issue td:first-child{{color:#854214;font-weight:650}}
th:first-child,td:first-child{{min-width:145px;white-space:normal;overflow-wrap:normal}}th:last-child,td:last-child{{min-width:360px;max-width:540px;white-space:normal;overflow-wrap:normal}}td:nth-child(2),td:nth-child(5){{min-width:180px}}
.ok td:first-child{{color:#255b48}}.note{{padding:16px;background:#eaf0f6;border-radius:8px}}footer{{margin-top:28px;color:#465d73}}
@media(max-width:650px){{main{{padding:22px 14px}}.cards{{grid-template-columns:repeat(2,1fr)}}h1{{font-size:25px}}}}
</style></head><body><main><div class="eyebrow">Financial Control Tower · local table review</div>
<h1>{title}</h1><p class="intro">A value-only comparison using explicit column mappings. Each line below is one numeric field for a unique pair, or one retained blocked/unpaired source row. Duplicate keys are never joined or summed.</p>
<div class="cards">{cards_html}</div><div class="links"><a href="differences.csv">Download comparison CSV</a><a href="comparison.json">Exact values (JSON)</a><a href="manifest.json">Source and file manifest</a></div>
<h2>Inputs and comparison rule</h2><ul>{inputs_html}</ul>
<p class="note">Unit: {escape(policy["unit"])}. Match when |right − left| ≤ max({escape(policy["absolute_tolerance"])}, {escape(policy["relative_tolerance"])} × max(|left|, |right|)). Relative tolerance is a fraction. Different currency text blocks subtraction. Keys are exact text, including leading zeros, case and surrounding spaces.</p>
<p>Row coverage: {summary["accounted_rows"]["left"]} / {summary["input_rows"]["left"]} left rows and {summary["accounted_rows"]["right"]} / {summary["input_rows"]["right"]} right rows retained. Blank row numbers and raw originals are in the JSON. CSV text is escaped against spreadsheet formula execution; import key/text columns as text to retain leading zeros.</p>
<h2>Field-level results</h2><div class="table-wrap"><table><thead><tr>{"".join(f"<th>{name}</th>" for name in headings)}</tr></thead><tbody>{"".join(row_html)}</tbody></table></div>
<footer>This is a technical comparison, not accounting reconciliation approval or financial-model validation. Sources are read locally; Excel formulas are not evaluated. Hashes detect changes but do not prove authenticity, independent origins or data correctness.</footer>
</main></body></html>"""
    path.write_text(document, encoding="utf-8")


def _publish_new(stage, output):
    """Atomically publish a directory without replacing even a concurrently created empty directory."""
    if sys.platform == "emscripten":
        # Browser comparisons use one serial Worker and its private synchronous
        # MEMFS. There is no second writer between this check and the rename.
        # Desktop platforms retain the native atomic no-replace implementation.
        if output.exists() or output.is_symlink():
            raise FileExistsError(str(output))
        os.rename(stage, output)
        return
    if sys.platform == "win32":
        os.rename(stage, output)  # Windows rename refuses any existing destination.
        return
    library = ctypes.CDLL(None, use_errno=True)
    if sys.platform == "darwin":
        function = library.renamex_np
        function.argtypes = [ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint]
        args = (os.fsencode(stage), os.fsencode(output), 4)  # RENAME_EXCL
    elif sys.platform.startswith("linux") and hasattr(library, "renameat2"):
        function = library.renameat2
        function.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p, ctypes.c_uint]
        args = (-100, os.fsencode(stage), -100, os.fsencode(output), 1)  # AT_FDCWD, RENAME_NOREPLACE
    else:
        raise OSError("this platform lacks a supported atomic no-replace directory rename")
    function.restype = ctypes.c_int
    if function(*args) != 0:
        error = ctypes.get_errno()
        raise OSError(error, os.strerror(error), str(output))


def run_comparison(left, right, output, *, absolute_tolerance="0", relative_tolerance="0", common_unit=None):
    """Read stable source snapshots, retain every input row, then publish a complete report."""
    absolute, relative = _validate_specs(left, right, absolute_tolerance, relative_tolerance, common_unit)
    output = Path(output).absolute()
    if output.exists() or output.is_symlink():
        raise FileExistsError(f"output already exists; choose a new directory: {output}")
    inputs, snapshots, tables = {}, {}, {}
    for side, spec in (("left", left), ("right", right)):
        content, identity, digest = _snapshot(spec.path)
        rows, metadata = _read_table(spec, content)
        snapshots[side] = (identity, digest)
        tables[side] = rows
        inputs[side] = {**asdict(spec), "path": str(Path(spec.path).absolute()), "sha256": digest, **metadata}
    potential_checks = (len(tables["left"]) + len(tables["right"])) * len(left.values)
    if potential_checks > MAX_FIELD_CHECKS:
        raise ValueError(
            "comparison exceeds 20,000 potential row-field entries before pairing; reduce source rows or mapped fields"
        )
    records = _compare_rows(tables["left"], tables["right"], left, right, absolute, relative)
    counts = dict(sorted(Counter(record["status"] for record in records).items()))
    input_rows = {side: len(rows) for side, rows in tables.items()}
    accounted = {
        side: len({record[f"{side}_row"] for record in records if record[f"{side}_row"] is not None})
        for side in ("left", "right")
    }
    if accounted != input_rows:
        raise RuntimeError("row accounting mismatch; no report published")
    result = {
        "schema_version": 1,
        "package_version": __version__,
        "inputs": inputs,
        "policy": {
            "absolute_tolerance": str(absolute),
            "relative_tolerance": str(relative),
            "unit": "explicit per-row currency columns"
            if left.currency
            else f"caller-declared common unit: {common_unit}",
            "difference": "right minus left; abs(diff) <= max(abs_tol, rel_tol * max(abs(left), abs(right)))",
            "keys": "exact text; no trimming, case conversion, numeric coercion, aggregation or deduplication",
            "formulas": "selected Excel formula cells are blocked even if a cached value exists",
            "csv_text": "apostrophe prefixes protect formula-like text; import identifiers as text; comparison.json retains originals",
            "scope": "offline technical comparison; not accounting or financial-model validation",
        },
        "summary": {
            "field_checks": counts,
            "input_rows": input_rows,
            "accounted_rows": accounted,
            "all_matched": counts == {"matched": len(records)},
        },
        "records": records,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=f".{output.name}-", dir=output.parent) as temporary:
        stage = Path(temporary) / "bundle"
        stage.mkdir()
        (stage / "comparison.json").write_text(
            json.dumps(result, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8"
        )
        _write_csv(stage / "differences.csv", records)
        _write_html(stage / "report.html", result)
        manifest = {
            "schema_version": 1,
            "package_version": __version__,
            "inputs": inputs,
            "policy": result["policy"],
            "summary": result["summary"],
            "sha256": {
                name: hashlib.sha256((stage / name).read_bytes()).hexdigest()
                for name in ("comparison.json", "differences.csv", "report.html")
            },
            "source_code_sha256": {
                path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                for path in (Path(__file__), Path(__file__).with_name("table_cli.py"))
            },
        }
        (stage / "manifest.json").write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
        for side, spec in (("left", left), ("right", right)):
            _content, identity, digest = _snapshot(spec.path)
            if snapshots[side] != (identity, digest):
                raise ValueError(f"{side} input changed during comparison; no report published")
        _publish_new(stage, output)
    return result
